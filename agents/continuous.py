"""Continuous debate loop: critics attack the best version, again and again.

This is NOT "build once, audit once, done". Every run:
  1. Critic attacks the current best version with deterministic checks
     that go BEYOND the audit gate (data labels, speaker notes, open
     evidence items).
  2. Defender checks for regression: a proposed fix must not break a
     check that already passed.
  3. Auditor re-scores the new version with office_builder.audit.
  4. The loop keeps the best version, logs the whole debate, and leaves
     an OPEN backlog. A temporary stop (2 rounds without improvement)
     is not a project close: the next run must reopen that backlog.

Usage:
    python -m agents.continuous examples/output runs/continuous-q3
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from office_builder.audit import audit_file  # noqa: E402
from office_builder.theme import load_tokens  # noqa: E402


def critic_issues(best_dir: Path) -> list[dict]:
    """Deterministic criticisms beyond the audit gate."""
    issues = []
    pptx_path = best_dir / "report.pptx"
    if pptx_path.exists():
        from pptx import Presentation

        prs = Presentation(str(pptx_path))
        for i, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.has_chart:
                    try:
                        has_labels = bool(shape.chart.plots[0].has_data_labels)
                    except Exception:
                        has_labels = False
                    if not has_labels:
                        issues.append({
                            "rule": "chart_data_labels", "severity": "MINOR",
                            "where": f"pptx slide {i} chart",
                            "measured": "has_data_labels=False",
                            "critic": "Người đọc phải ước lượng giá trị từ trục, thay vì đọc trực tiếp trên cột.",
                            "proposed_fix": "Bật data labels cho chart.",
                        })
            # Speaker notes are the presenter's safety net; a deck without
            # them is harder to present, even if it renders fine.
            try:
                notes = slide.notes_slide.placeholders[1].text_frame.text.strip() if slide.has_notes_slide else ""
            except Exception:
                notes = ""
            if not notes:
                issues.append({
                    "rule": "speaker_notes", "severity": "MINOR",
                    "where": f"pptx slide {i}",
                    "measured": "no speaker notes",
                    "critic": "Slide chỉ có chữ trên màn hình, người thuyết trình không có ghi chú dẫn dắt.",
                    "proposed_fix": "Thêm speaker notes tóm tắt ý chính của slide.",
                })
    xlsx_path = best_dir / "report.xlsx"
    if xlsx_path.exists():
        from openpyxl import load_workbook

        wb = load_workbook(str(xlsx_path))
        if "Dashboard" in wb.sheetnames:
            ws = wb["Dashboard"]
            for chart in ws._charts:
                if not getattr(chart, "dataLabels", None):
                    issues.append({
                        "rule": "xlsx_chart_data_labels", "severity": "MINOR",
                        "where": "xlsx Dashboard chart",
                        "measured": "dataLabels=None",
                        "critic": "Biểu đồ Excel cũng phải đọc được giá trị trực tiếp, không chỉ nhìn chiều cao cột.",
                        "proposed_fix": "Bật data labels cho chart Excel.",
                    })
    # Open evidence items can never be closed by code alone.
    issues.append({
        "rule": "powerpoint_playback_not_verified", "severity": "OPEN",
        "where": "pptx animation",
        "measured": "evidence rung = LibreOffice render, not interactive PowerPoint playback",
        "critic": "Animation mới chỉ qua cấu trúc + LibreOffice. Chưa ai bấm click thật trong PowerPoint.",
        "proposed_fix": "Giữ trong backlog mở; chỉ đóng khi có người trình chiếu click thật và ghi lại phiên bản PowerPoint.",
    })
    return issues


def apply_fixes(best_dir: Path, out_dir: Path, issues: list[dict]) -> None:
    """Apply only the fixable criticisms; OPEN items stay in the backlog."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for f in best_dir.iterdir():
        if f.is_file():
            shutil.copy2(f, out_dir / f.name)
    fixable = {i["rule"] for i in issues if i["severity"] != "OPEN"}
    if "chart_data_labels" in fixable or "speaker_notes" in fixable:
        from pptx import Presentation

        prs = Presentation(str(out_dir / "report.pptx"))
        for slide in prs.slides:
            if "chart_data_labels" in fixable:
                for shape in slide.shapes:
                    if shape.has_chart:
                        try:
                            shape.chart.plots[0].has_data_labels = True
                        except Exception:
                            pass
            if "speaker_notes" in fixable:
                try:
                    notes_slide = slide.notes_slide
                    tf = notes_slide.placeholders[1].text_frame
                    title = slide.shapes.title.text_frame.text if slide.shapes.title else "Slide"
                    tf.text = f"Ghi chú thuyết trình: {title}. Nói ý chính, chỉ vào số liệu, rồi chuyển slide ở click tiếp theo."
                except Exception:
                    pass
        prs.save(str(out_dir / "report.pptx"))
    if "xlsx_chart_data_labels" in fixable:
        from openpyxl import load_workbook
        from openpyxl.chart.label import DataLabelList

        wb = load_workbook(str(out_dir / "report.xlsx"))
        ws = wb["Dashboard"]
        for chart in ws._charts:
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
        wb.save(str(out_dir / "report.xlsx"))


def run_continuous(best_dir: str | Path, run_dir: str | Path, max_rounds: int = 3) -> dict:
    best_dir = Path(best_dir)
    run_dir = Path(run_dir)
    run_dir.mkdir(parents=True, exist_ok=True)
    tokens = load_tokens()
    log: list[str] = ["# Debate log — phản biện liên tục", ""]
    current_best = best_dir
    current_score = min(audit_file(p, tokens)["score"] for p in current_best.iterdir() if p.suffix in (".pptx", ".docx", ".xlsx"))
    log.append(f"Điểm khởi đầu (bản tốt nhất hiện tại): {current_score}")
    no_improvement = 0
    for round_no in range(1, max_rounds + 1):
        issues = critic_issues(current_best)
        fixable = [i for i in issues if i["severity"] != "OPEN"]
        open_items = [i for i in issues if i["severity"] == "OPEN"]
        log.append(f"\n## Vòng {round_no} — Critic")
        for i in issues:
            log.append(f"- [{i['severity']}] {i['rule']} tại {i['where']}: {i['measured']} — {i['critic']}")
        if not fixable:
            log.append("\nDefender: không còn phản biện nào sửa được bằng code trong vòng này.")
            log.append("Auditor: giữ nguyên bản tốt nhất. Backlog mở vẫn còn:")
            for i in open_items:
                log.append(f"- {i['rule']}: {i['proposed_fix']}")
            break
        out_dir = run_dir / f"round{round_no}"
        apply_fixes(current_best, out_dir, fixable)
        reports = [audit_file(p, tokens) for p in out_dir.iterdir() if p.suffix in (".pptx", ".docx", ".xlsx")]
        new_score = min(r["score"] for r in reports)
        blockers = sum(len(r["blockers"]) for r in reports)
        log.append(f"\nDefender: kiểm tra thoái lui — điểm mới {new_score}, blockers {blockers} (cũ {current_score}).")
        if blockers == 0 and new_score >= current_score:
            log.append(f"Auditor: CHẤP NHẬN cải tiến. Điểm {current_score} -> {new_score}. Bản mới trở thành bản tốt nhất.")
            current_best = out_dir
            if new_score > current_score:
                current_score = new_score
                no_improvement = 0
            else:
                no_improvement += 1
        else:
            log.append("Auditor: TỪ CHỐI — bản mới thoái lui. Giữ bản cũ.")
            no_improvement += 1
        log.append("Backlog mở sau vòng này:")
        for i in open_items:
            log.append(f"- {i['rule']}: {i['proposed_fix']}")
        if no_improvement >= 2:
            log.append("\nDừng tạm thời: 2 vòng liên tiếp không cải thiện điểm. Đây KHÔNG phải đóng dự án — lần chạy sau phải mở lại backlog trên.")
            break
    (run_dir / "debate_log.md").write_text("\n".join(log), encoding="utf-8")
    return {"best_dir": str(current_best), "best_score": current_score, "log": str(run_dir / "debate_log.md")}


if __name__ == "__main__":
    best, out = sys.argv[1], sys.argv[2]
    print(json.dumps(run_continuous(best, out), indent=2, ensure_ascii=False))
